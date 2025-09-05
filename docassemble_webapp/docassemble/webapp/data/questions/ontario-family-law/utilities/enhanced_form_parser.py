#!/usr/bin/env python3
"""
Enhanced Form Parser for Ontario Family Law Forms
Handles both modern and legacy Word form fields
"""

import csv
import json
import logging
from pathlib import Path
from typing import Dict, List, Optional, Any
from dataclasses import dataclass, asdict
from datetime import datetime
import zipfile
from lxml import etree

# Document parsing libraries
import PyPDF2
import pdfplumber
from docx import Document
import fitz  # PyMuPDF

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

@dataclass
class FormField:
    """Represents a single form field"""
    field_id: str
    field_name: str
    field_type: str  # text, date, checkbox, radio, number, currency, dropdown, etc.
    field_label: str
    required: bool = False
    validation_rules: str = ""
    max_length: Optional[int] = None
    options: Optional[List[str]] = None
    table_name: Optional[str] = None
    table_row: Optional[int] = None
    field_context: str = ""
    page_number: int = 0
    coordinates: Optional[Dict] = None
    help_text: Optional[str] = None

class EnhancedFormParser:
    """Enhanced parser that handles legacy Word form fields"""
    
    def __init__(self, file_path: str):
        self.file_path = Path(file_path)
        self.form_number = self.extract_form_number()
        self.fields = []
        
    def extract_form_number(self) -> str:
        """Extract form number from filename"""
        import re
        filename = self.file_path.name.lower()
        
        # Try to match form_XX pattern
        match = re.search(r'form[_\-](\d+(?:\.\d+)?[a-zA-Z]?)', filename)
        if match:
            return match.group(1).upper()
        
        # Try to match flr-XX pattern
        match = re.search(r'flr[_\-](\d+(?:\.\d+)?[a-zA-Z]?)', filename)
        if match:
            return match.group(1).upper()
        
        return ""
    
    def parse(self) -> List[FormField]:
        """Parse the form and extract fields"""
        ext = self.file_path.suffix.lower()
        
        if ext == '.pdf':
            return self.parse_pdf()
        elif ext in ['.docx', '.doc']:
            return self.parse_docx()
        else:
            logger.warning(f"Unsupported file type: {ext}")
            return []
    
    def parse_pdf(self) -> List[FormField]:
        """Parse PDF form - unchanged from original"""
        fields = []
        
        try:
            # Method 1: PyPDF2 for form fields
            fields.extend(self._parse_pdf_pypdf2())
        except Exception as e:
            logger.debug(f"PyPDF2 parsing failed: {e}")
        
        try:
            # Method 2: pdfplumber for text extraction
            fields.extend(self._parse_pdf_pdfplumber())
        except Exception as e:
            logger.debug(f"pdfplumber parsing failed: {e}")
        
        try:
            # Method 3: PyMuPDF for better field detection
            fields.extend(self._parse_pdf_pymupdf())
        except Exception as e:
            logger.debug(f"PyMuPDF parsing failed: {e}")
        
        # Deduplicate
        return self._deduplicate_fields(fields)
    
    def _parse_pdf_pypdf2(self) -> List[FormField]:
        """Parse PDF using PyPDF2"""
        fields = []
        
        with open(self.file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            
            if '/AcroForm' in reader.trailer['/Root']:
                form = reader.trailer['/Root']['/AcroForm']
                if '/Fields' in form:
                    for field_ref in form['/Fields']:
                        field_obj = field_ref.get_object()
                        field = self._extract_pypdf2_field(field_obj)
                        if field:
                            fields.append(field)
        
        return fields
    
    def _extract_pypdf2_field(self, field_obj) -> Optional[FormField]:
        """Extract field from PyPDF2 field object"""
        try:
            field_type = field_obj.get('/FT', '')
            field_name = field_obj.get('/T', '')
            field_flags = field_obj.get('/Ff', 0)
            
            # Map PDF field types to our types
            type_map = {
                '/Tx': 'text',
                '/Btn': 'checkbox' if field_flags & 0x10000 else 'radio',
                '/Ch': 'select',
                '/Sig': 'signature'
            }
            
            return FormField(
                field_id=f"pdf_{field_name}",
                field_name=field_name,
                field_type=type_map.get(field_type, 'text'),
                field_label=field_obj.get('/TU', field_name),  # Tooltip text
                required=bool(field_flags & 0x2),
                max_length=field_obj.get('/MaxLen', None),
                field_context="PDF Form Field"
            )
        except Exception as e:
            logger.debug(f"Failed to extract PyPDF2 field: {e}")
            return None
    
    def _parse_pdf_pdfplumber(self) -> List[FormField]:
        """Parse PDF using pdfplumber for text patterns"""
        fields = []
        
        with pdfplumber.open(self.file_path) as pdf:
            for page_num, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                fields.extend(self._extract_fields_from_text(text, page_num))
        
        return fields
    
    def _parse_pdf_pymupdf(self) -> List[FormField]:
        """Parse PDF using PyMuPDF"""
        fields = []
        
        doc = fitz.open(self.file_path)
        
        for page_num, page in enumerate(doc, 1):
            # Get form widgets
            for widget in page.widgets():
                field = self._extract_pymupdf_field(widget, page_num)
                if field:
                    fields.append(field)
            
            # Also extract text patterns
            text = page.get_text()
            fields.extend(self._extract_fields_from_text(text, page_num))
        
        doc.close()
        
        return fields
    
    def _extract_pymupdf_field(self, widget, page_num: int) -> Optional[FormField]:
        """Extract field from PyMuPDF widget"""
        try:
            field_type_map = {
                fitz.PDF_WIDGET_TYPE_TEXT: 'text',
                fitz.PDF_WIDGET_TYPE_CHECKBOX: 'checkbox',
                fitz.PDF_WIDGET_TYPE_RADIOBUTTON: 'radio',
                fitz.PDF_WIDGET_TYPE_LISTBOX: 'select',
                fitz.PDF_WIDGET_TYPE_COMBOBOX: 'select'
            }
            
            field_type = field_type_map.get(widget.field_type, 'text')
            
            return FormField(
                field_id=f"widget_{page_num}_{widget.field_name}",
                field_name=widget.field_name or f"field_{page_num}",
                field_type=field_type,
                field_label=widget.field_label or widget.field_name or "",
                required=widget.field_flags & 2 != 0,
                page_number=page_num,
                coordinates={
                    "x0": widget.rect.x0,
                    "y0": widget.rect.y0,
                    "x1": widget.rect.x1,
                    "y1": widget.rect.y1
                },
                field_context=f"Page {page_num}"
            )
        except Exception as e:
            logger.debug(f"Failed to extract PyMuPDF field: {e}")
            return None
    
    def parse_docx(self) -> List[FormField]:
        """Parse DOCX form - enhanced to handle legacy form fields"""
        fields = []
        
        try:
            # Method 1: Extract legacy form fields from XML
            fields.extend(self._extract_legacy_form_fields())
            
            # Method 2: Standard python-docx parsing
            doc = Document(self.file_path)
            
            # Extract from paragraphs
            for para_idx, paragraph in enumerate(doc.paragraphs):
                fields.extend(self._extract_fields_from_text(paragraph.text, para_idx))
            
            # Extract from tables
            for table_idx, table in enumerate(doc.tables):
                fields.extend(self._extract_docx_table_fields(table, table_idx))
            
        except Exception as e:
            logger.error(f"Failed to parse DOCX: {e}")
        
        # Deduplicate
        return self._deduplicate_fields(fields)
    
    def _extract_legacy_form_fields(self) -> List[FormField]:
        """Extract legacy Word form fields using XML parsing"""
        fields = []
        field_counter = {}
        
        try:
            # Open the DOCX as a ZIP file
            with zipfile.ZipFile(self.file_path, 'r') as docx:
                # Read the main document XML
                with docx.open('word/document.xml') as xml_file:
                    tree = etree.parse(xml_file)
                    root = tree.getroot()
                    
                    # Define namespaces
                    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                    
                    # Find all form field data elements
                    form_fields = root.xpath('//w:ffData', namespaces=ns)
                    
                    for ff_idx, ff in enumerate(form_fields):
                        # Extract field name
                        name_elem = ff.find('.//w:name', namespaces=ns)
                        if name_elem is not None:
                            field_name = name_elem.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', '')
                        else:
                            field_name = f'field_{ff_idx}'
                        
                        # Make unique field names
                        if field_name in field_counter:
                            field_counter[field_name] += 1
                            unique_name = f"{field_name}_{field_counter[field_name]}"
                        else:
                            field_counter[field_name] = 1
                            unique_name = field_name
                        
                        # Extract field type
                        field_type = 'text'
                        options = None
                        
                        if ff.find('.//w:checkBox', namespaces=ns) is not None:
                            field_type = 'checkbox'
                        elif ff.find('.//w:ddList', namespaces=ns) is not None:
                            field_type = 'dropdown'
                            # Extract dropdown options
                            list_entries = ff.findall('.//w:listEntry', namespaces=ns)
                            options = []
                            for entry in list_entries:
                                val = entry.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', '')
                                if val:
                                    options.append(val)
                            
                        # Extract help text
                        help_text = ''
                        help_elem = ff.find('.//w:helpText', namespaces=ns)
                        if help_elem is not None:
                            help_text = help_elem.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', '')
                        
                        # Extract status text (label)
                        status_text = ''
                        status_elem = ff.find('.//w:statusText', namespaces=ns)
                        if status_elem is not None:
                            status_text = status_elem.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', '')
                        
                        # Create human-readable label
                        label = status_text or self._humanize_field_name(field_name)
                        
                        fields.append(FormField(
                            field_id=f"legacy_{unique_name}",
                            field_name=self._sanitize_field_name(unique_name),
                            field_type=field_type,
                            field_label=label,
                            help_text=help_text,
                            options=options,
                            field_context="Legacy Form Field"
                        ))
                        
        except Exception as e:
            logger.debug(f"Failed to extract legacy form fields: {e}")
        
        return fields
    
    def _humanize_field_name(self, field_name: str) -> str:
        """Convert field name to human-readable label"""
        import re
        
        # Handle common patterns
        if field_name.lower() == 'courtfileno':
            return 'Court File Number'
        elif field_name.lower().startswith('text'):
            return f'Text Field {field_name[4:]}'.strip()
        elif field_name.lower().startswith('check'):
            return f'Checkbox {field_name[5:]}'.strip()
        elif field_name.lower().startswith('dropdown'):
            return f'Dropdown {field_name[8:]}'.strip()
        
        # General humanization
        label = re.sub(r'([A-Z])', r' \1', field_name)  # Add spaces before caps
        label = re.sub(r'_', ' ', label)  # Replace underscores
        label = re.sub(r'\s+', ' ', label)  # Clean up spaces
        label = label.strip().title()  # Title case
        
        return label
    
    def _extract_docx_table_fields(self, table, table_idx: int) -> List[FormField]:
        """Extract fields from DOCX table"""
        fields = []
        
        for row_idx, row in enumerate(table.rows):
            row_text = []
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    row_text.append(text)
            
            # Look for field patterns in row
            full_text = " ".join(row_text)
            if self._looks_like_field(full_text):
                fields.append(FormField(
                    field_id=f"table_{table_idx}_row_{row_idx}",
                    field_name=self._sanitize_field_name(full_text),
                    field_type=self._detect_field_type(full_text),
                    field_label=full_text[:100],
                    table_name=f"Table {table_idx + 1}",
                    table_row=row_idx,
                    field_context=f"Table {table_idx + 1}, Row {row_idx + 1}"
                ))
        
        return fields
    
    def _extract_fields_from_text(self, text: str, context_id: int) -> List[FormField]:
        """Extract fields from text using patterns"""
        fields = []
        
        if not text:
            return fields
        
        # Common field patterns
        patterns = [
            (r'(?:First|Last|Middle)\s+Name\s*[:_]?\s*_{3,}', 'text', 'name'),
            (r'Date\s+of\s+Birth\s*[:_]?\s*_{3,}', 'date', 'birthdate'),
            (r'Address\s*[:_]?\s*_{3,}', 'text', 'address'),
            (r'Phone\s*(?:Number)?\s*[:_]?\s*_{3,}', 'phone', 'phone'),
            (r'Email\s*(?:Address)?\s*[:_]?\s*_{3,}', 'email', 'email'),
            (r'(?:Court\s+)?File\s+(?:Number|No\.?)\s*[:_]?\s*_{3,}', 'text', 'file_number'),
            (r'□\s*(.+?)(?:\s*□|$)', 'checkbox', 'checkbox'),
            (r'○\s*(.+?)(?:\s*○|$)', 'radio', 'radio'),
            (r'\$\s*_{3,}', 'currency', 'amount'),
            (r'(\w+[\s\w]*?)\s*[:_]\s*_{3,}', 'text', 'generic'),
        ]
        
        for pattern, field_type, field_category in patterns:
            import re
            matches = re.finditer(pattern, text, re.IGNORECASE)
            for match in matches:
                label = match.group(1) if len(match.groups()) > 0 else match.group(0)
                label = label.strip(' :_')
                
                if label and len(label) > 2:
                    fields.append(FormField(
                        field_id=f"{field_category}_{context_id}_{len(fields)}",
                        field_name=self._sanitize_field_name(label),
                        field_type=field_type,
                        field_label=label,
                        field_context=f"Section {context_id}"
                    ))
        
        return fields
    
    def _looks_like_field(self, text: str) -> bool:
        """Check if text looks like a form field"""
        indicators = ['___', '[ ]', '( )', '□', '○', 'Name:', 'Date:', 'Address:', '$']
        text_lower = text.lower()
        return any(ind in text or ind.lower() in text_lower for ind in indicators)
    
    def _detect_field_type(self, text: str) -> str:
        """Detect field type from text"""
        text_lower = text.lower()
        
        if any(word in text_lower for word in ['date', 'when', 'born']):
            return 'date'
        elif any(word in text_lower for word in ['amount', 'value', 'income', '$']):
            return 'currency'
        elif any(word in text_lower for word in ['email', 'e-mail']):
            return 'email'
        elif any(word in text_lower for word in ['phone', 'telephone', 'fax']):
            return 'phone'
        elif any(word in text_lower for word in ['number', 'count', '#']):
            return 'number'
        elif '□' in text or '[ ]' in text:
            return 'checkbox'
        elif '○' in text or '( )' in text:
            return 'radio'
        
        return 'text'
    
    def _sanitize_field_name(self, text: str) -> str:
        """Create valid field name from text"""
        import re
        name = re.sub(r'[^a-zA-Z0-9_]', '_', text)
        name = re.sub(r'_+', '_', name)
        name = name.strip('_').lower()[:50]
        return name or "field"
    
    def _deduplicate_fields(self, fields: List[FormField]) -> List[FormField]:
        """Remove duplicate fields"""
        seen = set()
        unique = []
        
        for field in fields:
            key = (field.field_name, field.field_type, field.page_number or field.table_row or 0)
            if key not in seen:
                seen.add(key)
                unique.append(field)
        
        return unique

def parse_forms_with_enhanced_parser(csv_file: str = "workflow_output/forms_list.csv", output_dir: str = "workflow_output/enhanced_parsed_forms"):
    """Parse all forms using enhanced parser"""
    
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    
    results = []
    
    with open(csv_file, 'r') as f:
        reader = csv.DictReader(f)
        
        for row in reader:
            file_path = row['full_path']
            form_number = row['form_number']
            
            logger.info(f"Parsing Form {form_number}: {row['filename']}")
            
            try:
                parser = EnhancedFormParser(file_path)
                fields = parser.parse()
                
                if fields:
                    # Save as CSV
                    csv_file = output_path / f"form_{form_number}_fields.csv"
                    with open(csv_file, 'w', newline='') as cf:
                        writer = csv.DictWriter(cf, fieldnames=list(asdict(fields[0]).keys()))
                        writer.writeheader()
                        for field in fields:
                            writer.writerow(asdict(field))
                    
                    # Save as JSON
                    json_file = output_path / f"form_{form_number}_fields.json"
                    with open(json_file, 'w') as jf:
                        json.dump([asdict(f) for f in fields], jf, indent=2)
                    
                    logger.info(f"  ✓ Extracted {len(fields)} fields")
                    results.append({
                        "form_number": form_number,
                        "filename": row['filename'],
                        "fields_count": len(fields),
                        "status": "success"
                    })
                else:
                    logger.warning(f"  ⚠ No fields extracted")
                    results.append({
                        "form_number": form_number,
                        "filename": row['filename'],
                        "fields_count": 0,
                        "status": "no_fields"
                    })
                    
            except Exception as e:
                logger.error(f"  ✗ Error: {e}")
                results.append({
                    "form_number": form_number,
                    "filename": row['filename'],
                    "fields_count": 0,
                    "status": "error",
                    "error": str(e)
                })
    
    # Save summary
    summary = {
        "timestamp": datetime.now().isoformat(),
        "total_forms": len(results),
        "successful": len([r for r in results if r["status"] == "success"]),
        "no_fields": len([r for r in results if r["status"] == "no_fields"]),
        "errors": len([r for r in results if r["status"] == "error"]),
        "results": results
    }
    
    summary_file = output_path / "enhanced_parsing_summary.json"
    with open(summary_file, 'w') as f:
        json.dump(summary, f, indent=2)
    
    logger.info(f"\nParsing complete. Summary saved to {summary_file}")
    
    return results

if __name__ == "__main__":
    parse_forms_with_enhanced_parser()