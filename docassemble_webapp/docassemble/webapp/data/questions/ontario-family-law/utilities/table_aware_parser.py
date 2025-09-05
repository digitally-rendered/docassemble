#!/usr/bin/env python3
"""
Table-Aware Parser for Ontario Family Law Forms
Extracts fields from both tables and paragraphs in DOCX/PDF files
"""

import os
import json
import csv
import logging
from pathlib import Path
from typing import Dict, List, Optional, Any, Tuple
from dataclasses import dataclass, asdict
from datetime import datetime
import re

# Import parsers
import PyPDF2
import pdfplumber
from docx import Document
import fitz  # PyMuPDF

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

@dataclass
class ExtractedField:
    """Represents an extracted field from a form"""
    field_id: str
    label: str
    field_type: str
    page_number: int
    table_number: Optional[int]
    row_number: Optional[int]
    cell_number: Optional[int]
    required: bool
    options: List[str]
    metadata: Dict

    def to_dict(self) -> Dict:
        return asdict(self)

class TableAwareParser:
    """Parser that extracts fields from tables and paragraphs"""
    
    # Common field labels in Ontario Family Law forms
    FIELD_LABELS = [
        'Court File Number', 'Court Name', 'Court office address',
        'Full legal name', 'Name', 'First name', 'Last name',
        'Address', 'Phone', 'Fax', 'Email',
        'Date of Birth', 'Birthdate', 'Age',
        'Resident in', 'municipality', 'province',
        'Gender', 'Male', 'Female', 'Another gender',
        'Applicant', 'Respondent', 'Lawyer',
        'Date of marriage', 'Date of separation',
        'Number of children', 'Child\'s name',
        'Income', 'Expenses', 'Assets', 'Debts',
        'Monthly amount', 'Annual income',
        'Yes', 'No', 'Check the box',
        'Provide details', 'Specify', 'Explain',
        'I am asking for', 'I want the court to',
        'Support', 'Custody', 'Access', 'Property',
        'Divorce', 'Other claims'
    ]
    
    def __init__(self):
        self.fields = []
        self.field_counter = 0
        
    def parse_form(self, file_path: str) -> Dict:
        """Parse a form and extract all fields"""
        logger.info(f"Parsing form: {Path(file_path).name}")
        
        self.fields = []
        self.field_counter = 0
        
        file_ext = Path(file_path).suffix.lower()
        
        if file_ext in ['.docx', '.doc']:
            self._parse_docx(file_path)
        elif file_ext == '.pdf':
            self._parse_pdf(file_path)
        else:
            logger.warning(f"Unsupported file type: {file_ext}")
            
        return self._create_result()
    
    def _parse_docx(self, file_path: str):
        """Parse DOCX file and extract fields from tables and paragraphs"""
        try:
            doc = Document(file_path)
            
            # Extract fields from tables
            for table_idx, table in enumerate(doc.tables):
                self._extract_fields_from_table(table, table_idx, 1)
            
            # Extract fields from paragraphs
            for para_idx, para in enumerate(doc.paragraphs):
                self._extract_fields_from_text(para.text, 1, para_idx)
                
        except Exception as e:
            logger.error(f"Failed to parse DOCX: {e}")
    
    def _parse_pdf(self, file_path: str):
        """Parse PDF file and extract fields"""
        try:
            with pdfplumber.open(file_path) as pdf:
                for page_num, page in enumerate(pdf.pages, 1):
                    # Extract tables
                    tables = page.extract_tables()
                    for table_idx, table in enumerate(tables):
                        self._extract_fields_from_pdf_table(table, table_idx, page_num)
                    
                    # Extract text
                    text = page.extract_text()
                    if text:
                        self._extract_fields_from_text(text, page_num)
                        
        except Exception as e:
            logger.error(f"Failed to parse PDF: {e}")
    
    def _extract_fields_from_table(self, table, table_idx: int, page_num: int):
        """Extract fields from a DOCX table"""
        for row_idx, row in enumerate(table.rows):
            row_text = []
            for cell_idx, cell in enumerate(row.cells):
                cell_text = cell.text.strip()
                if cell_text:
                    # Check if this looks like a field
                    field_info = self._identify_field(cell_text, page_num)
                    if field_info:
                        field_info['table_number'] = table_idx
                        field_info['row_number'] = row_idx
                        field_info['cell_number'] = cell_idx
                        self.fields.append(field_info)
    
    def _extract_fields_from_pdf_table(self, table: List, table_idx: int, page_num: int):
        """Extract fields from a PDF table"""
        for row_idx, row in enumerate(table):
            if not row:
                continue
            for cell_idx, cell in enumerate(row):
                if cell and cell.strip():
                    # Check if this looks like a field
                    field_info = self._identify_field(cell, page_num)
                    if field_info:
                        field_info['table_number'] = table_idx
                        field_info['row_number'] = row_idx
                        field_info['cell_number'] = cell_idx
                        self.fields.append(field_info)
    
    def _extract_fields_from_text(self, text: str, page_num: int, line_num: int = None):
        """Extract fields from plain text"""
        if not text:
            return
            
        lines = text.split('\n')
        for idx, line in enumerate(lines):
            if line.strip():
                field_info = self._identify_field(line, page_num)
                if field_info:
                    field_info['line_number'] = line_num or idx
                    self.fields.append(field_info)
    
    def _identify_field(self, text: str, page_num: int) -> Optional[Dict]:
        """Identify if text represents a field"""
        text = text.strip()
        if not text:
            return None
        
        # Check for common field patterns
        for label in self.FIELD_LABELS:
            if label.lower() in text.lower():
                return self._create_field(text, label, page_num)
        
        # Check for field indicators
        if ':' in text and len(text) < 100:
            # This might be a field label
            parts = text.split(':', 1)
            label = parts[0].strip()
            if label and len(label) < 50:
                return self._create_field(text, label, page_num)
        
        # Check for checkbox patterns
        if any(marker in text for marker in ['☐', '□', '[ ]', '( )']):
            return self._create_field(text, text, page_num, field_type='checkbox')
        
        # Check for Yes/No options
        if text.lower() in ['yes', 'no'] and len(text) < 10:
            return self._create_field(text, text, page_num, field_type='checkbox')
        
        return None
    
    def _create_field(self, text: str, label: str, page_num: int, field_type: str = None) -> Dict:
        """Create a field entry"""
        self.field_counter += 1
        
        # Determine field type if not provided
        if not field_type:
            field_type = self._determine_field_type(label)
        
        # Check if required
        required = self._is_required_field(label)
        
        # Extract options for checkboxes/radio buttons
        options = self._extract_options(text) if field_type in ['checkbox', 'radio'] else []
        
        field = ExtractedField(
            field_id=f"field_{self.field_counter:04d}",
            label=label,
            field_type=field_type,
            page_number=page_num,
            table_number=None,
            row_number=None,
            cell_number=None,
            required=required,
            options=options,
            metadata={'original_text': text[:200]}
        )
        
        return field.to_dict()
    
    def _determine_field_type(self, label: str) -> str:
        """Determine field type based on label"""
        label_lower = label.lower()
        
        if any(word in label_lower for word in ['date', 'birth', 'marriage', 'separation']):
            return 'date'
        elif 'email' in label_lower:
            return 'email'
        elif any(word in label_lower for word in ['phone', 'fax', 'telephone']):
            return 'phone'
        elif any(word in label_lower for word in ['amount', 'income', 'expense', 'cost', '$']):
            return 'currency'
        elif any(word in label_lower for word in ['number of', 'how many']):
            return 'number'
        elif any(word in label_lower for word in ['yes', 'no', 'check']):
            return 'checkbox'
        elif any(word in label_lower for word in ['address', 'street', 'city', 'postal']):
            return 'address'
        elif 'gender' in label_lower:
            return 'radio'
        else:
            return 'text'
    
    def _is_required_field(self, label: str) -> bool:
        """Determine if field is required"""
        label_lower = label.lower()
        required_fields = [
            'court file number', 'court name', 'applicant', 'respondent',
            'full legal name', 'name', 'date of birth'
        ]
        return any(req in label_lower for req in required_fields)
    
    def _extract_options(self, text: str) -> List[str]:
        """Extract options from checkbox/radio text"""
        options = []
        
        # Look for Yes/No pattern
        if 'yes' in text.lower() and 'no' in text.lower():
            return ['Yes', 'No']
        
        # Look for Male/Female/Other pattern
        if 'male' in text.lower() and 'female' in text.lower():
            options = ['Male', 'Female']
            if 'another' in text.lower() or 'other' in text.lower():
                options.append('Another gender')
            return options
        
        return options
    
    def _create_result(self) -> Dict:
        """Create parsing result"""
        return {
            'total_fields': len(self.fields),
            'fields': self.fields,
            'field_types': self._get_field_type_stats(),
            'required_fields': [f for f in self.fields if f.get('required')],
            'optional_fields': [f for f in self.fields if not f.get('required')]
        }
    
    def _get_field_type_stats(self) -> Dict:
        """Get statistics on field types"""
        stats = {}
        for field in self.fields:
            field_type = field.get('field_type', 'unknown')
            stats[field_type] = stats.get(field_type, 0) + 1
        return stats
    
    def save_to_csv(self, result: Dict, output_path: str):
        """Save parsed fields to CSV"""
        csv_path = Path(output_path)
        csv_path.parent.mkdir(parents=True, exist_ok=True)
        
        if not result['fields']:
            logger.warning(f"No fields to save to {csv_path}")
            return
        
        with open(csv_path, 'w', newline='', encoding='utf-8') as f:
            fieldnames = [
                'field_id', 'label', 'field_type', 'page_number',
                'table_number', 'row_number', 'cell_number',
                'required', 'options'
            ]
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            
            for field in result['fields']:
                row = {
                    'field_id': field['field_id'],
                    'label': field['label'],
                    'field_type': field['field_type'],
                    'page_number': field['page_number'],
                    'table_number': field.get('table_number', ''),
                    'row_number': field.get('row_number', ''),
                    'cell_number': field.get('cell_number', ''),
                    'required': field['required'],
                    'options': '|'.join(field.get('options', []))
                }
                writer.writerow(row)
        
        logger.info(f"Saved {len(result['fields'])} fields to {csv_path}")
    
    def save_to_json(self, result: Dict, output_path: str):
        """Save parsed fields to JSON"""
        json_path = Path(output_path)
        json_path.parent.mkdir(parents=True, exist_ok=True)
        
        with open(json_path, 'w', encoding='utf-8') as f:
            json.dump(result, f, indent=2)
        
        logger.info(f"Saved parsing result to {json_path}")

def test_parser():
    """Test the table-aware parser"""
    parser = TableAwareParser()
    
    # Test with a sample form
    test_form = "workflow_output/parsed_data/downloads/form_8_-_application_general_flr-8-jun25-en.docx"
    
    if os.path.exists(test_form):
        result = parser.parse_form(test_form)
        
        print(f"\n=== PARSING RESULTS ===")
        print(f"Total fields found: {result['total_fields']}")
        print(f"\nField types: {result['field_types']}")
        print(f"Required fields: {len(result['required_fields'])}")
        print(f"Optional fields: {len(result['optional_fields'])}")
        
        print(f"\n=== SAMPLE FIELDS ===")
        for field in result['fields'][:10]:
            print(f"- {field['label']} ({field['field_type']})")
            if field.get('options'):
                print(f"  Options: {field['options']}")
        
        # Save results
        parser.save_to_csv(result, "test_form_8_fields.csv")
        parser.save_to_json(result, "test_form_8_fields.json")
    else:
        print(f"Test form not found: {test_form}")

if __name__ == "__main__":
    test_parser()